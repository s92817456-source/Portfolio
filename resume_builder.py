"""resume_builder.py — Claude rewrites the resume for a JD, then exports ATS-safe PDF/DOCX.
Rule baked into the prompt: NEVER invent experience. Missing skills the candidate has no evidence for
go to 'gaps_not_in_resume' (to learn / honestly add) instead of being faked in.
"""
from __future__ import annotations
import io
import re

from llm import call_claude, parse_json

_SYS = """You are an expert resume writer and ATS specialist. You rewrite resumes to match a job description
TRUTHFULLY. Never invent employers, titles, degrees, dates, tools or numbers. Reply with JSON only."""


def rewrite_resume(resume_text: str, job_text: str, missing_must: list[str], missing_nice: list[str]) -> dict:
    prompt = f"""Rewrite this resume for the job below to score 80%+ in ATS tools (Jobscan style).

RULES
- Use ONLY facts present in the resume. Reword, reorder, and use the JD's exact phrasing for keywords.
- Missing keywords: {missing_must[:15]} (must) / {missing_nice[:8]} (nice).
  Add one ONLY where the resume shows evidence (e.g. a project that used it). Otherwise put it in "gaps_not_in_resume".
- Where a metric would help but the resume has none, write a placeholder like [X%] or [N users] for the candidate to fill.
- Bullets start with strong action verbs. Plain single-column content, standard headings.
- Summary: 3 lines, includes the job title and top JD keywords.

Return JSON:
{{"name": str, "contact": "email | phone | linkedin | city",
 "summary": str,
 "skills": [{{"group": str, "items": [str]}}],
 "experience": [{{"title": str, "company": str, "dates": str, "bullets": [str]}}],
 "projects": [{{"name": str, "bullets": [str]}}],
 "education": [{{"degree": str, "school": str, "dates": str}}],
 "certifications": [str],
 "keyword_plan": [{{"keyword": str, "where_to_add": str, "suggested_line": str}}],
 "gaps_not_in_resume": [str],
 "changes": [str]}}

JOB DESCRIPTION:
{job_text[:3500]}

RESUME:
{resume_text[:6000]}"""
    return parse_json(call_claude(prompt, _SYS, max_tokens=3800))


def resume_to_text(d: dict) -> str:
    out = [d.get("name", ""), d.get("contact", ""), "SUMMARY", d.get("summary", ""), "SKILLS"]
    out += [f"{g.get('group', '')}: {', '.join(g.get('items', []))}" for g in d.get("skills", [])]
    out.append("EXPERIENCE")
    for e in d.get("experience", []):
        out += [f"{e.get('title', '')} - {e.get('company', '')} ({e.get('dates', '')})"] + [f"- {b}" for b in e.get("bullets", [])]
    if d.get("projects"):
        out.append("PROJECTS")
        for p in d["projects"]:
            out += [p.get("name", "")] + [f"- {b}" for b in p.get("bullets", [])]
    out.append("EDUCATION")
    out += [f"{e.get('degree', '')} - {e.get('school', '')} ({e.get('dates', '')})" for e in d.get("education", [])]
    if d.get("certifications"):
        out += ["CERTIFICATIONS"] + [f"- {c}" for c in d["certifications"]]
    return "\n".join(out)


def _s(t) -> str:
    t = "" if t is None else str(t)
    for a, b in {"–": "-", "—": "-", "•": "-", "’": "'", "‘": "'", "“": '"', "”": '"',
                 "…": "...", "→": "->", "\u00a0": " "}.items():
        t = t.replace(a, b)
    t = t.encode("latin-1", "replace").decode("latin-1")
    return re.sub(r"(\S{60})(?=\S)", r"\1 ", t)


def build_resume_pdf(d: dict) -> io.BytesIO:
    from fpdf import FPDF
    pdf = FPDF(format="A4")
    pdf.set_margins(15, 14, 15)
    pdf.set_auto_page_break(True, 14)
    pdf.add_page()

    def put(txt, size=10, style="", h=5, align="L"):
        pdf.set_font("Helvetica", style, size)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, h, _s(txt), align=align, new_x="LMARGIN", new_y="NEXT")

    def head(t):
        pdf.ln(2)
        put(t.upper(), 11, "B", 6)
        pdf.line(15, pdf.get_y(), 195, pdf.get_y())
        pdf.ln(1)

    put(d.get("name", ""), 18, "B", 8, "C")
    put(d.get("contact", ""), 9, "", 5, "C")
    head("Summary")
    put(d.get("summary", ""))
    head("Skills")
    for g in d.get("skills", []):
        put(f"{g.get('group', '')}: {', '.join(g.get('items', []))}", 10)
    head("Experience")
    for e in d.get("experience", []):
        put(f"{e.get('title', '')} | {e.get('company', '')} | {e.get('dates', '')}", 10, "B")
        for b in e.get("bullets", []):
            put(f"- {b}")
        pdf.ln(1)
    if d.get("projects"):
        head("Projects")
        for p in d["projects"]:
            put(p.get("name", ""), 10, "B")
            for b in p.get("bullets", []):
                put(f"- {b}")
    head("Education")
    for e in d.get("education", []):
        put(f"{e.get('degree', '')} | {e.get('school', '')} | {e.get('dates', '')}")
    if d.get("certifications"):
        head("Certifications")
        for c in d["certifications"]:
            put(f"- {c}")
    return io.BytesIO(bytes(pdf.output()))


def build_resume_docx(d: dict) -> io.BytesIO:
    from docx import Document          # pip install python-docx
    doc = Document()
    doc.add_heading(d.get("name", ""), 0)
    doc.add_paragraph(d.get("contact", ""))

    def sec(title):
        doc.add_heading(title, 1)

    sec("Summary"); doc.add_paragraph(d.get("summary", ""))
    sec("Skills")
    for g in d.get("skills", []):
        doc.add_paragraph(f"{g.get('group', '')}: {', '.join(g.get('items', []))}")
    sec("Experience")
    for e in d.get("experience", []):
        p = doc.add_paragraph(); p.add_run(f"{e.get('title', '')} | {e.get('company', '')} | {e.get('dates', '')}").bold = True
        for b in e.get("bullets", []):
            doc.add_paragraph(b, style="List Bullet")
    if d.get("projects"):
        sec("Projects")
        for pr in d["projects"]:
            p = doc.add_paragraph(); p.add_run(pr.get("name", "")).bold = True
            for b in pr.get("bullets", []):
                doc.add_paragraph(b, style="List Bullet")
    sec("Education")
    for e in d.get("education", []):
        doc.add_paragraph(f"{e.get('degree', '')} | {e.get('school', '')} | {e.get('dates', '')}")
    if d.get("certifications"):
        sec("Certifications")
        for c in d["certifications"]:
            doc.add_paragraph(c, style="List Bullet")
    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    return buf
